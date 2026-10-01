"""Low-level helpers for reading and editing Word tables with python-docx.

We deliberately work on the raw table XML for rows/cells (instead of table.rows[i].cells)
because python-docx resolves vertically merged cells to the cell ABOVE, which would make us
overwrite the wrong row. Here every cell is addressed exactly where it physically is.
"""
import copy

from docx.oxml.ns import qn
from docx.table import _Cell

from .common import clean_ws


def tr_grid_cells(tr):
    """List of <w:tc> elements for each grid column of a row (spanned cells repeat)."""
    cells = []
    for tc in tr.tc_lst:
        span = 1
        vals = tc.xpath("./w:tcPr/w:gridSpan/@w:val")
        if vals:
            try:
                span = max(1, int(vals[0]))
            except ValueError:
                span = 1
        cells.extend([tc] * span)
    return cells


def is_vmerge_continue(tc) -> bool:
    """True for the lower cells of a vertically merged region (they hold no content)."""
    nodes = tc.xpath("./w:tcPr/w:vMerge")
    if not nodes:
        return False
    return nodes[0].get(qn("w:val")) in (None, "continue")


def has_merge(tc) -> bool:
    return bool(tc.xpath("./w:tcPr/w:vMerge")) or bool(tc.xpath("./w:tcPr/w:gridSpan"))


def tc_text(tc, table) -> str:
    """Plain text of a cell (paragraphs joined by newlines)."""
    if is_vmerge_continue(tc):
        return ""
    return _Cell(tc, table).text.strip()


def row_texts(tr, table) -> list:
    """Cell texts for every grid column of a row."""
    cache = {}
    out = []
    for tc in tr_grid_cells(tr):
        key = id(tc)
        if key not in cache:
            cache[key] = tc_text(tc, table)
        out.append(cache[key])
    return out


def paragraph_text(p_element) -> str:
    return "".join(t.text or "" for t in p_element.iter(qn("w:t"))).strip()


def iter_tables_with_heading(doc):
    """Yield (table, heading) for each top-level table; heading = paragraph right before it."""
    tables = doc.tables
    idx, last = 0, ""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = paragraph_text(child)
            if text:
                last = text
        elif child.tag == qn("w:tbl"):
            if idx < len(tables):
                yield tables[idx], last
            idx += 1
            last = ""


def is_blank_row(texts, date_col, day_col=None) -> bool:
    """A row with no date and nothing else (a pre-printed 'day' name is allowed)."""
    if date_col is not None and date_col < len(texts) and texts[date_col].strip():
        return False
    for i, t in enumerate(texts):
        if i == date_col or i == day_col:
            continue
        if t.strip():
            return False
    return True


def same_text(a, b) -> bool:
    return clean_ws(a) == clean_ws(b)


def _remove_non_content(tr):
    """Remove bookmarks/comment markers from a cloned row (their ids must stay unique)."""
    for tag in ("w:bookmarkStart", "w:bookmarkEnd", "w:commentRangeStart", "w:commentRangeEnd"):
        for el in tr.xpath(".//" + tag):
            el.getparent().remove(el)


def clear_tc(tc):
    """Empty a cell but keep its first paragraph and first run (so formatting survives)."""
    paragraphs = tc.findall(qn("w:p"))
    for extra in paragraphs[1:]:
        tc.remove(extra)
    if not paragraphs:
        return
    p = paragraphs[0]
    runs = p.findall(qn("w:r"))
    keep = runs[0] if runs else None
    for child in list(p):
        if child.tag == qn("w:pPr") or child is keep:
            continue
        p.remove(child)
    if keep is not None:
        for child in list(keep):
            if child.tag != qn("w:rPr"):
                keep.remove(child)


def clone_row_empty(tr):
    """Deep-copy a row and empty its cells. Merged cells keep their structure."""
    new = copy.deepcopy(tr)
    _remove_non_content(new)
    for tc in new.tc_lst:
        clear_tc(tc)
    return new


def set_tc_text(tc, table, text):
    """Replace a cell's text while keeping paragraph and run formatting."""
    cell = _Cell(tc, table)
    paragraphs = cell.paragraphs
    para = paragraphs[0] if paragraphs else cell.add_paragraph()
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    runs = para.runs
    if runs:
        first = runs[0]
        first.text = text  # python-docx turns "\n" into a line break
        for child in list(para._p):
            if child.tag == qn("w:pPr") or child is first._r:
                continue
            para._p.remove(child)
    else:
        run = para.add_run(text)
        mark = para._p.xpath("./w:pPr/w:rPr")
        if mark:  # paragraph-mark formatting is the best guess for an empty cell's font
            run._r.insert(0, copy.deepcopy(mark[0]))
