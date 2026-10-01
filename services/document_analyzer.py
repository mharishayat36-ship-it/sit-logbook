"""Analyses a Word document: finds tables, header rows, columns and the date column."""
from docx import Document

from .common import auto_columns, clean_ws, match_field, parse_date
from .word_utils import (has_merge, is_blank_row, iter_tables_with_heading, row_texts,
                         tr_grid_cells)

MAX_HEADER_SEARCH_ROWS = 6


def _guess_name(field_cols) -> str:
    fields = set(field_cols)
    detail = {"task", "project", "domain", "tools", "learning"}
    if fields & detail:
        return "Daily Progress"
    if fields & {"check_in", "check_out"}:
        return "Attendance"
    return ""


def _find_header_row(all_texts) -> int:
    best_row, best_score = 0, 0
    for i, texts in enumerate(all_texts[:MAX_HEADER_SEARCH_ROWS]):
        score = sum(1 for t in texts if match_field(t)[0])
        if score > best_score:
            best_row, best_score = i, score
    return best_row


def analyze_table(table, index, heading="", default_year=None) -> dict:
    trs = table._tbl.tr_lst
    all_texts = [row_texts(tr, table) for tr in trs]
    header_row = _find_header_row(all_texts) if all_texts else 0
    headers = [clean_ws(t) for t in all_texts[header_row]] if all_texts else []
    field_cols = auto_columns(headers)

    # A merged title row above the header (e.g. "WEEK 1 ATTENDANCE") works as a heading.
    if not heading and header_row > 0:
        uniq = {t for t in all_texts[0] if t}
        if len(uniq) == 1:
            heading = clean_ws(next(iter(uniq)))
    heading = clean_ws(heading)
    if len(heading) > 60:
        heading = ""

    data_rows = all_texts[header_row + 1:]
    date_col = field_cols.get("date")
    if date_col is None and data_rows:  # look for a column that is mostly dates
        width = max(len(r) for r in data_rows)
        for c in range(width):
            vals = [r[c] for r in data_rows if c < len(r) and r[c].strip()]
            if vals and sum(parse_date(v, default_year) is not None for v in vals) >= 0.6 * len(vals):
                date_col = c
                break

    day_col = field_cols.get("day")
    blank = sum(is_blank_row(r, date_col, day_col) for r in data_rows)
    empty = sum(not any(c.strip() for c in r) for r in data_rows)
    samples = [[t[:60] for t in r] for r in data_rows if any(c.strip() for c in r)][:3]
    merged = any(has_merge(tc) for tr in trs for tc in tr.tc_lst)

    guess = _guess_name(field_cols)
    name = heading or guess or f"Table {index + 1}"
    notes = []
    if merged:
        notes.append("Contains merged cells. Merged cells are preserved; content inside "
                     "continuation cells of vertical merges is never changed.")
    if date_col is None:
        notes.append("No date column was detected - rows cannot be matched by date.")
    if not trs:
        notes.append("The table has no rows.")
    return {
        "index": index, "number": index + 1, "name": name, "heading": heading, "guess": guess,
        "rows": len(trs), "cols": max((len(tr_grid_cells(tr)) for tr in trs), default=0),
        "header_row": header_row, "columns": headers,
        "date_col": date_col,
        "date_col_name": headers[date_col] if date_col is not None and date_col < len(headers) else "",
        "field_columns": field_cols,
        "data_rows": len(data_rows), "empty_rows": empty, "blank_rows": blank,
        "samples": samples, "has_merged": merged, "notes": notes,
    }


def analyze_loaded(doc, default_year=None) -> list:
    """Analyse an already-opened python-docx Document."""
    return [analyze_table(t, i, h, default_year)
            for i, (t, h) in enumerate(iter_tables_with_heading(doc))]


def analyze_document(path, default_year=None) -> dict:
    doc = Document(path)
    tables = analyze_loaded(doc, default_year)
    warnings = []
    if not tables:
        warnings.append("This document contains no tables. Only table-based logbooks are supported.")
    section_info = {"sections": len(doc.sections),
                    "header_text": clean_ws(" ".join(p.text for p in doc.sections[0].header.paragraphs)),
                    "footer_text": clean_ws(" ".join(p.text for p in doc.sections[0].footer.paragraphs))}
    return {"table_count": len(tables), "tables": tables, "warnings": warnings, **section_info}
