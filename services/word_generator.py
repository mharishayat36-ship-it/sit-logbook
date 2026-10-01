"""Writes records into a copy of the Word template.

Two phases, so nothing is changed without confirmation:

    build_plan()  - read-only. Works out every cell that would change (the preview).
    apply_plan()  - loads the ORIGINAL file again, applies the plan and saves a NEW file.

The original template is never modified.
"""
from datetime import date

from docx import Document

from .common import (DEFAULT_DATE_STYLE, FIELD_LABELS, clean_ws, day_name, detect_date_style,
                     format_date, long_date, normalize_time, parse_date, parse_day_name)
from .document_analyzer import analyze_loaded
from .field_mapper import build_targets, mapped_fields_of, target_matches_name
from .word_utils import (clone_row_empty, is_blank_row, is_vmerge_continue, row_texts,
                         set_tc_text, tr_grid_cells)

MODES = {
    "update": "Update existing rows only",
    "add": "Add missing rows",
    "strict": "Strict (error if the date is missing)",
}
CORE_UNMAPPED_ERROR = {"task", "project", "domain", "tools", "learning", "check_in", "check_out"}


# ---------------------------------------------------------------- helpers
def merge_by_date(flat_records, table_records=None) -> dict:
    """{iso: {field: value}} from flat records (later ones win)."""
    out = {}
    for rec in flat_records:
        out.setdefault(rec["date"], {}).update({k: v for k, v in rec.items() if v != ""})
    return out


def _date_rows(table, info, date_col, default_year):
    """{iso: first row index} for every row whose date cell can be read."""
    found = {}
    trs = table._tbl.tr_lst
    for i in range(info["header_row"] + 1, len(trs)):
        texts = row_texts(trs[i], table)
        if date_col < len(texts):
            d = parse_date(texts[date_col], default_year)
            if d and d.isoformat() not in found:
                found[d.isoformat()] = i
    return found


def _nearest_table(d, table_indexes, lookups):
    """For a date missing from a group of tables (e.g. one table per week) pick the table
    whose existing dates are closest to it (distance 0 = the date falls inside its range)."""
    best, best_dist = table_indexes[-1], None
    for ti in table_indexes:
        dates = [date.fromisoformat(k) for k in lookups[ti]]
        if not dates:
            dist = 10 ** 6
        elif min(dates) <= d <= max(dates):
            dist = 0
        else:
            dist = min(abs((d - min(dates)).days), abs((d - max(dates)).days))
        if best_dist is None or dist < best_dist:
            best, best_dist = ti, dist
    return best


def _blank_rows(table, info, date_col):
    trs = table._tbl.tr_lst
    day_col = info["field_columns"].get("day")
    return [i for i in range(info["header_row"] + 1, len(trs))
            if is_blank_row(row_texts(trs[i], table), date_col, day_col)]


def _table_date_style(table, info, date_col):
    trs = table._tbl.tr_lst
    for i in range(info["header_row"] + 1, len(trs)):
        texts = row_texts(trs[i], table)
        if date_col < len(texts):
            style = detect_date_style(texts[date_col])
            if style:
                return style
    return DEFAULT_DATE_STYLE


def _same_value(field, old, new) -> bool:
    if field in ("check_in", "check_out"):
        return normalize_time(old)[0] == normalize_time(new)[0]
    if field == "day":
        return parse_day_name(old) == parse_day_name(new) and parse_day_name(new) is not None
    if field == "date":
        a, b = parse_date(old), parse_date(new)
        return a is not None and a == b
    return clean_ws(old) == clean_ws(new)


# ---------------------------------------------------------------- phase 1: plan
def build_plan(doc_path, flat_records, table_records, mappings, mode, default_year=None) -> dict:
    """Describe (without changing anything) what inserting the records would do."""
    doc = Document(doc_path)
    tables = analyze_loaded(doc, default_year)
    word_tables = doc.tables
    table_records = table_records or {}
    targets, issues = build_targets(tables, mappings, extra_names=list(table_records.keys()))
    plan = {"mode": mode, "mode_label": MODES.get(mode, mode), "changes": [], "issues": list(issues),
            "targets": [], "stats": {"update": 0, "add": 0, "skip": 0, "error": 0, "nochange": 0}}

    if not tables:
        plan["issues"].append({"level": "error", "text": "The document contains no tables."})
        return plan
    if not targets:
        plan["issues"].append({"level": "error", "text":
                               "No Word table matches the field mappings. Action: open Field Mapping "
                               "and use 'Auto-detect from document'."})
        return plan

    # fields that have no Word column at all
    everything = list(flat_records) + [r for rows in table_records.values() for r in rows]
    used = {f for r in everything for f, v in r.items() if v and f in FIELD_LABELS and f != "date"}
    for f in sorted(used - mapped_fields_of(targets)):
        plan["issues"].append({"level": "error" if f in CORE_UNMAPPED_ERROR else "warning", "text":
                               f"Field '{f}': no matching Word column was found. "
                               f"Action: configure field mapping. (The field is not written.)"})

    for target in targets:
        # merge flat records with the records meant for this specific table
        combined = merge_by_date(flat_records)
        for key, rows in table_records.items():
            if target_matches_name(key, target, tables):
                for rec in rows:
                    combined.setdefault(rec["date"], {}).update({k: v for k, v in rec.items() if v != ""})
        if not combined:
            continue
        plan["targets"].append({"name": target["name"],
                                "tables": [tables[i]["name"] for i in target["table_indexes"]],
                                "source": target["source"]})
        # per-table lookups
        lookups, blanks, styles = {}, {}, {}
        for ti in target["table_indexes"]:
            dc = target["date_cols"][ti]
            lookups[ti] = _date_rows(word_tables[ti], tables[ti], dc, default_year)
            blanks[ti] = _blank_rows(word_tables[ti], tables[ti], dc)
            styles[ti] = _table_date_style(word_tables[ti], tables[ti], dc)

        for iso in sorted(combined):
            rec = combined[iso]
            d = date.fromisoformat(iso)
            change = {"table": target["name"], "date": iso, "date_label": long_date(d),
                      "cells": [], "message": "", "action": "", "table_index": None,
                      "row_number": None, "date_col": None}

            # Does this record have anything for this target besides date/day?
            relevant = [f for f in rec if f not in ("date", "day") and rec[f]
                        and any(f in cols for cols in target["columns"].values())]
            if not relevant:
                continue

            home = next((ti for ti in target["table_indexes"] if iso in lookups[ti]), None)
            if home is not None:
                action, ti, row = "update", home, lookups[home][iso]
            elif mode == "add":
                near = _nearest_table(d, target["table_indexes"], lookups)
                if blanks[near]:
                    action, ti, row = "fill_blank", near, blanks[near].pop(0)
                else:
                    action, ti, row = "add", near, None
            elif mode == "strict":
                change.update(action="error", message=
                              f"Date {long_date(d)} does not exist in table '{target['name']}'. "
                              f"Strict mode: record not applied.")
                plan["changes"].append(change)
                plan["stats"]["error"] += 1
                continue
            else:
                change.update(action="skip", message=
                              f"Date {long_date(d)} was not found in table '{target['name']}'. "
                              f"Update mode: record skipped.")
                plan["changes"].append(change)
                plan["stats"]["skip"] += 1
                continue

            cols = target["columns"][ti]
            table, info = word_tables[ti], tables[ti]
            existing = row_texts(table._tbl.tr_lst[row], table) if row is not None else []
            grid = tr_grid_cells(table._tbl.tr_lst[row]) if row is not None else []
            change.update(action=action, table_index=ti, date_col=target["date_cols"][ti],
                          row_number=(row + 1) if row is not None else None)
            seen_tc = set()
            for field in [f for f in FIELD_LABELS if f in cols]:
                value = rec.get(field, "")
                if field == "date":
                    if action == "update":
                        continue
                    value = format_date(d, styles[ti])
                elif not value:
                    continue
                col = cols[field]
                old = existing[col] if col < len(existing) else ""
                if grid and col < len(grid):
                    if is_vmerge_continue(grid[col]):
                        plan["issues"].append({"level": "warning", "text":
                            f"{long_date(d)} / '{info['columns'][col]}' is inside a vertically merged "
                            f"cell and was left unchanged."})
                        continue
                    if id(grid[col]) in seen_tc:
                        plan["issues"].append({"level": "warning", "text":
                            f"{long_date(d)}: two fields map to the same merged cell "
                            f"('{info['columns'][col]}'); only the first was used."})
                        continue
                    seen_tc.add(id(grid[col]))
                header = info["columns"][col] if col < len(info["columns"]) else f"column {col + 1}"
                change["cells"].append({"field": field, "column": header, "col": col, "old": old,
                                        "new": value, "changed": not _same_value(field, old, value)
                                        or action != "update"})
            if action == "update" and not any(c["changed"] for c in change["cells"]):
                change["action"] = "nochange"
                change["message"] = "Already up to date."
                plan["stats"]["nochange"] += 1
            else:
                plan["stats"]["update" if action == "update" else "add"] += 1
                if action == "add":
                    change["message"] = "A new row will be inserted in date order."
                elif action == "fill_blank":
                    change["message"] = "An empty row will be filled."
            plan["changes"].append(change)
    return plan


def plan_has_work(plan) -> bool:
    return any(c["action"] in ("update", "add", "fill_blank") for c in plan["changes"])


# ---------------------------------------------------------------- phase 2: apply
def _insert_row(table, info, date_col, iso, default_year):
    """Insert an empty clone of a neighbouring row at the correct chronological place."""
    trs = table._tbl.tr_lst
    start = info["header_row"] + 1
    target = date.fromisoformat(iso)
    dated = []
    for i in range(start, len(trs)):
        texts = row_texts(trs[i], table)
        d = parse_date(texts[date_col], default_year) if date_col < len(texts) else None
        if d:
            dated.append((i, d))
    after_idx = None
    for i, d in dated:
        if d > target:
            after_idx = i
            break
    if after_idx is not None:
        ref = trs[after_idx]
        new = clone_row_empty(ref)
        ref.addprevious(new)
    elif dated:
        ref = trs[dated[-1][0]]
        new = clone_row_empty(ref)
        ref.addnext(new)
    else:
        ref = trs[-1]
        new = clone_row_empty(ref)
        ref.addnext(new)
    return new


def apply_plan(src_path, dst_path, plan, default_year=None) -> dict:
    """Apply a plan to a fresh load of `src_path` and save the result to `dst_path`."""
    doc = Document(src_path)
    tables = analyze_loaded(doc, default_year)
    word_tables = doc.tables
    applied, warnings = 0, []
    for ch in plan["changes"]:
        if ch["action"] not in ("update", "add", "fill_blank"):
            continue
        ti = ch["table_index"]
        table, info, dc = word_tables[ti], tables[ti], ch["date_col"]
        trs = table._tbl.tr_lst
        tr = None
        if ch["action"] == "update":
            row = _date_rows(table, info, dc, default_year).get(ch["date"])
            tr = trs[row] if row is not None else None
        elif ch["action"] == "fill_blank":
            blanks = _blank_rows(table, info, dc)
            tr = trs[blanks[0]] if blanks else None
        if tr is None and ch["action"] in ("add", "fill_blank"):
            tr = _insert_row(table, info, dc, ch["date"], default_year)
        if tr is None:
            warnings.append(f"{ch['date']}: the row could not be located; skipped.")
            continue
        grid = tr_grid_cells(tr)
        done = set()
        for cell in ch["cells"]:
            if not cell["changed"]:
                continue
            col = cell["col"]
            if col >= len(grid):
                warnings.append(f"{ch['date']}: column {col + 1} does not exist in the row; skipped.")
                continue
            tc = grid[col]
            if is_vmerge_continue(tc) or id(tc) in done:
                continue
            done.add(id(tc))
            set_tc_text(tc, table, cell["new"])
        applied += 1
    doc.save(dst_path)
    return {"applied": applied, "warnings": warnings}
