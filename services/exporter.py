"""Export records as JSON, CSV or Excel (.xlsx)."""
import csv
import io
import json

from .common import CORE_FIELDS, FIELDS


def export_fields(records) -> list:
    """The 9 standard fields, plus 'notes' only when some record uses it."""
    fields = list(CORE_FIELDS)
    if any(r.get("notes") for r in records):
        fields.append("notes")
    return fields


def clean_records(records) -> list:
    fields = export_fields(records)
    return [{f: str(r.get(f, "") or "") for f in fields} for r in records]


def to_json(records) -> str:
    return json.dumps(clean_records(records), indent=4, ensure_ascii=False)


def to_csv(records) -> str:
    fields = export_fields(records)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=fields, lineterminator="\r\n")
    writer.writeheader()
    for r in clean_records(records):
        writer.writerow(r)
    return buf.getvalue()


def to_xlsx_bytes(records) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    fields = export_fields(records)
    wb = Workbook()
    ws = wb.active
    ws.title = "Logbook"
    ws.append(fields)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5D8A")
    for r in clean_records(records):
        ws.append([r[f] for f in fields])
    widths = {"date": 12, "day": 12, "check_in": 11, "check_out": 11, "task": 48, "project": 22,
              "domain": 22, "tools": 24, "learning": 48, "notes": 30}
    for i, f in enumerate(fields, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(f, 20)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
