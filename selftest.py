"""End-to-end self test. Uses a temporary database and folders - your real data is not touched.

Run:  python selftest.py
"""
import hashlib
import io
import json
import os
import re
import shutil
import sys
import tempfile

from docx import Document

import create_sample_template
from app import create_app

results = []


def check(name, condition, detail=""):
    results.append((name, bool(condition)))
    print(("PASS  " if condition else "FAIL  ") + name + (f"   <- {detail}" if detail and not condition else ""))


def sha(path):
    with open(path, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


def main():
    tmp = tempfile.mkdtemp(prefix="logbook_test_")
    app = create_app({
        "DATABASE_URI": "sqlite:///" + os.path.join(tmp, "test.db").replace("\\", "/"),
        "UPLOAD_FOLDER": os.path.join(tmp, "uploads"), "GENERATED_FOLDER": os.path.join(tmp, "generated"),
        "EXPORT_FOLDER": os.path.join(tmp, "exports"), "TESTING": True, "SECRET_KEY": "test"})
    c = app.test_client()

    # ---- every page loads
    for url in ["/", "/daily", "/weekly", "/json/import", "/extract", "/analyzer", "/mappings",
                "/generate", "/entries", "/history", "/history?tab=extractions", "/history?tab=documents"]:
        check(f"GET {url}", c.get(url).status_code == 200)
    check("404 page", c.get("/nope").status_code == 404)

    # ---- upload
    sample = create_sample_template.build()
    sample_hash = sha(sample)
    r = c.post("/documents/upload", data={"document": (open(sample, "rb"), "My Logbook.docx"), "next": "analyzer"},
               content_type="multipart/form-data", follow_redirects=True)
    check("upload sample docx", b"Uploaded" in r.data and b"Attendance" in r.data and b"Daily Progress" in r.data)
    stored = os.listdir(app.config["UPLOAD_FOLDER"])
    check("stored with safe unique name", len(stored) == 1 and stored[0].endswith("My_Logbook.docx"), str(stored))
    for bad, label in [((io.BytesIO(b"hello"), "evil.exe"), "reject .exe"),
                       ((io.BytesIO(b"not a zip"), "fake.docx"), "reject fake .docx"),
                       ((io.BytesIO(b"PK\x03\x04junk"), "bad.docx"), "reject corrupt .docx")]:
        r = c.post("/documents/upload", data={"document": bad}, content_type="multipart/form-data", follow_redirects=True)
        check(label, b"Uploaded '" not in r.data and len(os.listdir(app.config["UPLOAD_FOLDER"])) == 1)

    # ---- JSON -> Word (multiple records with a mismatch and an invalid date)
    payload = [
        {"date": "2026-09-01", "day": "Tuesday", "check_in": "09:00 AM", "check_out": "05:00 PM",
         "task": "Worked on ball physics.", "project": "Cup Toss Game", "domain": "Physics Programming",
         "tools": "Unity, C#, Rigidbody", "learning": "Learned Rigidbody physics."},
        {"date": "2026-09-02", "day": "Monday", "check_in": "9am", "check_out": "17:00",
         "task": "Collision detection.", "project": "Cup Toss Game", "domain": "Collision", "tools": "Unity", "learning": "Colliders"},
        {"date": "2026-09-31", "task": "bad date"},
        {"date": "2026-09-03", "task": "T", "mystery": "x"},
    ]
    r = c.post("/json/import", data={"json_text": json.dumps(payload)}, follow_redirects=False)
    check("JSON import redirects to preview", r.status_code == 302)
    preview_url = r.headers["Location"]
    draft_id = int(re.search(r"/(\d+)$", preview_url).group(1))
    r = c.get(preview_url)
    html = r.data.decode()
    check("preview shows invalid calendar date", "Invalid calendar date" in html)
    check("preview shows day mismatch warning", "is Wednesday, but the supplied day is Monday" in html)
    check("preview shows unknown field warning", "Unknown field" in html)
    check("preview offers correct/keep/cancel", "Correct automatically" in html and "Keep the supplied value" in html and "Cancel import" in html)

    r = c.post(f"/generate/{draft_id}/preview", data={"mode": "update", "day_policy": "correct", "doc_id": "1"})
    html = r.data.decode()
    check("changes preview renders", r.status_code == 200 and "Preview of changes" in html and "Confirm and Generate" in html)
    check("change stats render as numbers", "built-in method" not in html and "bound method" not in html)
    check("changes preview lists new values", "Worked on ball physics." in html and "EMPTY" in html)
    r = c.post(f"/generate/{draft_id}/confirm", data={"mode": "update", "day_policy": "correct", "doc_id": "1", "also_save": "1"})
    check("generate succeeds", r.status_code == 200 and b"Document generated" in r.data)
    gen_dir = app.config["GENERATED_FOLDER"]
    generated = os.listdir(gen_dir)
    check("generated file named *_Completed.docx", len(generated) == 1 and generated[0].endswith("_Completed.docx"), str(generated))
    check("original template untouched", sha(sample) == sample_hash)
    doc = Document(os.path.join(gen_dir, generated[0]))
    t_att, t_prog = doc.tables
    row = lambda t, i: [cell.text for cell in t.rows[i].cells]
    check("attendance row filled (Sep 1)", row(t_att, 2) == ["01 Sep 2026", "Tuesday", "09:00 AM", "05:00 PM"], str(row(t_att, 2)))
    check("day mismatch auto-corrected (Sep 2 = Wednesday)", row(t_att, 3)[1] == "Wednesday" and row(t_att, 3)[2] == "09:00 AM" and row(t_att, 3)[3] == "05:00 PM", str(row(t_att, 3)))
    check("progress row filled (Sep 1)", row(t_prog, 2)[1:] == ["Worked on ball physics.", "Cup Toss Game", "Physics Programming", "Unity, C#, Rigidbody", "Learned Rigidbody physics."], str(row(t_prog, 2)))
    check("untouched rows stay empty", row(t_prog, 5)[1] == "" and row(t_att, 5)[2] == "")
    check("header/footer preserved", doc.sections[0].header.paragraphs[0].text == "Internship Training Logbook" and doc.sections[0].footer.paragraphs[0].text == "Student Industrial Training")
    check("row count unchanged in update mode", len(t_att.rows) == 23 and len(t_prog.rows) == 23)

    # ---- strict + add modes with a date outside the template
    far = [{"date": "2026-10-01", "day": "Thursday", "check_in": "09:00 AM", "check_out": "05:00 PM", "task": "Far away task", "project": "P", "domain": "D", "tools": "T", "learning": "L"}]
    r = c.post("/json/import", data={"json_text": json.dumps(far)})
    far_id = int(re.search(r"/(\d+)$", r.headers["Location"]).group(1))
    r = c.post(f"/generate/{far_id}/preview", data={"mode": "strict", "day_policy": "correct", "doc_id": "1"})
    check("strict mode reports error, no confirm button", b"Strict mode: record not applied" in r.data and b"Confirm and Generate" not in r.data)
    r = c.post(f"/generate/{far_id}/preview", data={"mode": "update", "day_policy": "correct", "doc_id": "1"})
    check("update mode skips missing date", b"Update mode: record skipped" in r.data and b"Confirm and Generate" not in r.data)
    r = c.post(f"/generate/{far_id}/preview", data={"mode": "add", "day_policy": "correct", "doc_id": "1"})
    check("add mode offers new rows", b"Add new row" in r.data and b"Confirm and Generate" in r.data)
    r = c.post(f"/generate/{far_id}/confirm", data={"mode": "add", "day_policy": "correct", "doc_id": "1"})
    gen2 = sorted(os.listdir(gen_dir))
    check("second generation gets its own file (no overwrite)", len(gen2) == 2, str(gen2))
    newest = max((os.path.join(gen_dir, g) for g in gen2), key=os.path.getmtime)
    d2 = Document(newest)
    check("added row inserted in date order after Sep 26", len(d2.tables[0].rows) == 24 and d2.tables[0].rows[-1].cells[0].text == "01 Oct 2026" and d2.tables[0].rows[-1].cells[1].text == "Thursday")

    # ---- table-specific JSON
    ts = {"tables": {"Attendance": [{"date": "2026-09-04", "day": "Friday", "check_in": "09:00 AM", "check_out": "05:00 PM"}],
                     "Daily Progress": [{"date": "2026-09-04", "task": "Only progress", "project": "X"}]}}
    r = c.post("/json/import", data={"json_text": json.dumps(ts)})
    ts_id = int(re.search(r"/(\d+)$", r.headers["Location"]).group(1))
    r = c.post(f"/generate/{ts_id}/confirm", data={"mode": "update", "day_policy": "correct", "doc_id": "1"})
    newest = max((os.path.join(gen_dir, g) for g in os.listdir(gen_dir)), key=os.path.getmtime)
    d3 = Document(newest)
    check("table-specific JSON fills the right tables", d3.tables[0].rows[5].cells[2].text == "09:00 AM" and d3.tables[1].rows[5].cells[1].text == "Only progress" and d3.tables[1].rows[5].cells[2].text == "X")
    bad_tbl = {"tables": {"Nonexistent": [{"date": "2026-09-04", "task": "x"}]}}
    r = c.post("/json/import", data={"json_text": json.dumps(bad_tbl)}, follow_redirects=True)
    check("unknown target table is an error", b"was not found in the Word document" in r.data)

    # ---- JSON syntax error / single object / repair
    r = c.post("/json/import", data={"json_text": '{"date": "2026-09-28",'}, follow_redirects=True)
    check("syntax error reported with position", b"JSON syntax error at line" in r.data)
    r = c.post("/json/import", data={"json_text": json.dumps({"date": "2026-09-28", "task": "x"})})
    check("single object accepted", r.status_code == 302)
    r = c.post("/json/import", data={"json_text": '```json\n{"date": "2026-09-28", "task": "x",}\n```'})
    check("fenced JSON with trailing comma repaired", r.status_code == 302)

    # ---- save to DB path
    r = c.post(f"/json/{draft_id}/save-db", data={"day_policy": "correct"})
    check("save JSON to database", r.status_code == 200 and b"Saved to database" in r.data)
    r = c.get("/entries")
    check("entries list shows saved records", b"01 Sep 2026" in r.data and b"Worked on ball physics." in r.data)
    r = c.get("/history")
    check("import history lists imports", b"Import #" in r.data)

    # ---- Word -> JSON extraction
    r = c.post("/extract/run", data={"doc_id": "1", "skip_empty": "1"})
    check("extract on template with 2 filled days works", r.status_code == 302, r.data[:200].decode(errors="ignore"))
    # extract from the generated (completed) document: upload it
    comp = os.path.join(gen_dir, generated[0])
    r = c.post("/documents/upload", data={"document": (open(comp, "rb"), "Completed.docx"), "next": "extract"},
               content_type="multipart/form-data", follow_redirects=True)
    r = c.post("/extract/run", data={"doc_id": "2", "skip_empty": "1"})
    check("extract redirects to preview", r.status_code == 302)
    ex_url = r.headers["Location"]
    ex_id = int(re.search(r"/(\d+)$", ex_url).group(1))
    r = c.get(ex_url)
    html = r.data.decode()
    check("extraction preview lists records and warnings", "Extracted records: 2" in html and "Learning field is empty" in html or "field is empty" in html, html[:300])
    check("combined attendance+progress", "Worked on ball physics." in html and "09:00 AM" in html)
    # editing API
    ex_records = json.loads(c.get(f"/extract/{ex_id}/export/json").data)
    check("export JSON has 9 fields", list(ex_records[0].keys()) == ["date", "day", "check_in", "check_out", "task", "project", "domain", "tools", "learning"], str(ex_records[0].keys()))
    check("extracted record correct", ex_records[0]["date"] == "2026-09-01" and ex_records[0]["tools"] == "Unity, C#, Rigidbody" and ex_records[0]["check_out"] == "05:00 PM")
    edited = [dict(r_, task="Edited task") for r_ in ex_records]
    r = c.post(f"/draft/{ex_id}/save", json={"records": edited})
    check("save edited records", r.status_code == 200 and r.get_json()["ok"])
    edited.append({"date": "2026-02-30", "task": "bad"})
    r = c.post(f"/draft/{ex_id}/save", json={"records": edited})
    check("save rejects invalid row", r.status_code == 400 and r.get_json()["rows"][-1]["status"] == "error")
    r = c.get(f"/extract/{ex_id}/export/json")
    check("edits persisted", json.loads(r.data)[0]["task"] == "Edited task")
    r = c.get(f"/extract/{ex_id}/export/csv")
    check("CSV export", r.status_code == 200 and r.data.decode("utf-8-sig").splitlines()[0] == "date,day,check_in,check_out,task,project,domain,tools,learning")
    r = c.get(f"/extract/{ex_id}/export/xlsx")
    check("Excel export", r.status_code == 200 and r.data[:2] == b"PK")
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.data)).active
    check("Excel contents", ws["A1"].value == "date" and ws["A2"].value == "2026-09-01")

    # ---- round trip
    r = c.get(f"/roundtrip/{draft_id}?doc_id=1")
    html = r.data.decode()
    check("round-trip report", r.status_code == 200 and "Round-Trip Validation" in html and "Date preserved" in html and "Word formatting" in html, html[:400])
    check("round-trip says all preserved", "All data preserved" in html)

    # ---- weekly report
    r = c.post("/weekly", data={"week": "2", "project": "3D Cube Runner", "action": "distribute",
                                 "progress": "Learned Unity Editor basics, GameObjects and Components, player movement, Rigidbody physics, camera follow, collision detection, scoring and UI.",
                                 "check_in": "09:00 AM", "check_out": "05:00 PM"})
    check("weekly distribution redirects", r.status_code == 302)
    wk_url = r.headers["Location"]
    wk_id = int(re.search(r"/(\d+)$", wk_url).group(1))
    r = c.get(wk_url)
    html = r.data.decode()
    check("week 2 has 6 daily rows incl. Saturday", html.count('data-field="date"') >= 6 and "Saturday" in html, str(html.count('data-field="date"')))
    wk = c.application.test_client()
    r = c.post("/weekly", data={"week": "3", "project": "P", "action": "distribute", "progress": "A, B, C, D, E, F, G"})
    html = c.get(r.headers["Location"]).data.decode()
    check("week 3 has 5 days and no Saturday/Sunday", "Saturday" not in html.replace("Monday&ndash;Saturday", "") and "Sunday" not in html and html.count('data-field="date"') >= 5)
    r = c.post(f"/draft/{wk_id}/to-database")
    check("weekly saved to database", r.status_code == 302)
    r = c.post("/weekly", data={"week": "abc", "progress": "x", "action": "save"})
    check("weekly invalid week handled", b"Week must be a whole number" in r.data)

    # ---- manual entry
    r = c.post("/daily", data={"date": "2026-09-14", "day": "Friday", "task": "Manual task", "check_in": "9", "check_out": "5pm"}, follow_redirects=True)
    check("manual entry saved + day corrected", b"Manual task" in r.data and b"corrected to Monday" in r.data, r.data.decode()[:300])
    r = c.post("/daily", data={"date": "2026-13-45"}, follow_redirects=True)
    check("manual entry invalid date rejected", b"Saved the entry" not in r.data)
    r = c.get("/api/day-info?date=2026-09-27")
    check("Sunday not a working day", r.get_json()["working"] is False)
    r = c.get("/api/day-info?date=2026-09-26")
    check("Saturday week 4 working", r.get_json()["working"] is True)
    r = c.get("/api/day-info?date=2026-09-19")
    check("Saturday week 3 not working", r.get_json()["working"] is False)

    # ---- generate from database
    r = c.post("/generate", data={"doc_id": "1", "mode": "add"})
    check("generate-from-database preview", r.status_code == 200 and b"Preview of changes" in r.data)

    # ---- mappings
    r = c.get("/mappings")
    check("mappings page lists defaults", b"Daily Progress" in r.data and b"Check Out" in r.data)
    r = c.post("/mappings/auto", follow_redirects=True)
    check("auto-detect mappings", b"Detected" in r.data)
    r = c.post("/mappings", data={"json_field": ["bogus"], "word_table": ["T"], "word_column": ["C"]}, follow_redirects=True)
    check("mapping validation", b"not a known JSON field" in r.data)
    c.post("/mappings/reset")

    # ---- security
    check("download blocks traversal", c.get("/download/generated/..%2f..%2fconfig.py").status_code == 404)
    check("download blocks bad kind", c.get("/download/uploads/x.docx").status_code == 404)
    r = c.post("/json/import", data={"json_text": json.dumps({"date": "2026-09-28", "task": "<script>alert(1)</script>"})})
    html = c.get(r.headers["Location"]).data.decode()
    check("HTML is escaped in previews", "<script>alert(1)</script>" not in html and "&lt;script&gt;" in html)

    shutil.rmtree(tmp, ignore_errors=True)
    failed = [n for n, ok in results if not ok]
    print(f"\n{len(results) - len(failed)}/{len(results)} checks passed")
    if failed:
        print("FAILED:", *failed, sep="\n  - ")
        sys.exit(1)


if __name__ == "__main__":
    main()

