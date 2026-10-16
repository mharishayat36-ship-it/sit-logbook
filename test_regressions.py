"""Regression coverage for saving, exports, validation, and form protection."""
import csv
import io
import json
import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

from app import create_app
from create_sample_template import build
from models import DailyEntry, Draft, UploadedDocument, WeeklyReport, db_session
from routes.helpers import create_draft, get_schedule, upsert_entries
from services import exporter
from services.file_storage import reserve_output
from services.json_processor import parse_json_text


class RegressionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='logbook_regression_')
        self.root = Path(self.tmp.name)
        self.app = create_app({
            'DATABASE_URI': 'sqlite:///:memory:', 'TESTING': True,
            'CSRF_ENABLED': False, 'SECRET_KEY': 'regression-test',
            'UPLOAD_FOLDER': str(self.root / 'uploads'),
            'GENERATED_FOLDER': str(self.root / 'generated'),
            'EXPORT_FOLDER': str(self.root / 'exports'),
        })
        self.client = self.app.test_client()
        build(output_file=str(self.root / 'uploads' / 'template.docx'))
        with self.app.app_context():
            self.engine = db_session.get_bind()
            doc = UploadedDocument(original_name='Logbook.docx', stored_name='template.docx', is_active=True)
            db_session.add(doc)
            db_session.commit()
            self.doc_id = doc.id

    def tearDown(self):
        db_session.remove()
        self.engine.dispose()
        self.tmp.cleanup()

    def import_json(self, payload):
        response = self.client.post('/json/import', data={'json_text': json.dumps(payload)})
        self.assertEqual(response.status_code, 302)
        return int(response.headers['Location'].rsplit('/', 1)[1])

    def make_draft(self, records, kind='weekly'):
        with self.app.app_context():
            return create_draft(kind, records, {'week': 1}).id

    def test_table_specific_records_merge_for_new_date(self):
        draft = self.import_json({'tables': {
            'Attendance': [{'date': '2026-09-01', 'check_in': '9am'}],
            'Daily Progress': [{'date': '2026-09-01', 'task': 'Built a prototype'}],
        }})
        response = self.client.post(f'/json/{draft}/save-db')
        self.assertEqual(response.status_code, 200)
        with self.app.app_context():
            self.assertEqual(db_session.query(DailyEntry).count(), 1)
            entry = db_session.query(DailyEntry).one()
            self.assertEqual(entry.check_in, '09:00 AM')
            self.assertEqual(entry.task, 'Built a prototype')

    def test_merge_preserves_existing_nonempty_values(self):
        with self.app.app_context():
            schedule = get_schedule()
            self.assertEqual(upsert_entries([{'date': '2026-09-01', 'task': 'Original', 'tools': 'Unity'}], schedule), (1, 0))
            self.assertEqual(upsert_entries([{'date': '2026-09-01', 'task': '', 'project': 'Game'}], schedule), (0, 1))
            entry = db_session.query(DailyEntry).one()
            self.assertEqual((entry.task, entry.tools, entry.project), ('Original', 'Unity', 'Game'))

    def test_invalid_weekly_time_is_rejected_before_report_saved(self):
        response = self.client.post('/weekly', data={'week': '2', 'progress': 'Prototype', 'check_in': 'invalid-time'})
        self.assertIn(b'Invalid check-in time', response.data)
        with self.app.app_context():
            self.assertEqual(db_session.query(WeeklyReport).count(), 0)
            self.assertEqual(db_session.query(Draft).count(), 0)

    def test_invalid_draft_cannot_save_or_generate(self):
        draft = self.make_draft([{'date': '2026-09-01', 'task': 'Prototype', 'check_in': 'invalid-time'}])
        response = self.client.post(f'/draft/{draft}/to-database', follow_redirects=True)
        self.assertIn(b'Fix the records', response.data)
        response = self.client.post(f'/generate/{draft}/confirm', data={'doc_id': self.doc_id}, follow_redirects=True)
        self.assertIn(b'Fix the records', response.data)
        self.assertEqual(list((self.root / 'generated').iterdir()), [])
        with self.app.app_context():
            self.assertEqual(db_session.query(DailyEntry).count(), 0)

    def test_validation_is_atomic_for_mixed_valid_invalid_records(self):
        with self.app.app_context():
            with self.assertRaises(ValueError):
                upsert_entries([{'date': '2026-09-01', 'task': 'Good'}, {'date': '2026-09-02', 'check_in': 'bad'}], get_schedule())
            self.assertEqual(db_session.query(DailyEntry).count(), 0)

    def test_three_generations_keep_distinct_files_and_content(self):
        for n in range(3):
            draft = self.make_draft([{'date': '2026-09-01', 'task': f'Version {n}'}])
            response = self.client.post(f'/generate/{draft}/confirm', data={'doc_id': self.doc_id})
            self.assertEqual(response.status_code, 200)
        files = list((self.root / 'generated').glob('*.docx'))
        self.assertEqual(len(files), 3)
        tasks = {Document(str(path)).tables[1].rows[2].cells[1].text for path in files}
        self.assertEqual(tasks, {'Version 0', 'Version 1', 'Version 2'})
        self.assertEqual(Document(str(self.root / 'uploads' / 'template.docx')).tables[1].rows[2].cells[1].text, '')

    def test_cancel_does_not_generate(self):
        draft = self.make_draft([{'date': '2026-09-01', 'task': 'Example'}])
        self.client.post(f'/generate/{draft}/confirm', data={'doc_id': self.doc_id, 'day_policy': 'cancel'})
        self.assertEqual(list((self.root / 'generated').iterdir()), [])

    def test_exports_do_not_overwrite_in_same_second(self):
        draft = self.make_draft([{'date': '2026-09-01', 'task': 'Example'}])
        with patch('routes.extract_routes.timestamp', return_value='same_second'):
            for _ in range(2):
                response = self.client.get(f'/extract/{draft}/export/json')
                self.assertEqual(response.status_code, 200)
                response.close()
        self.assertEqual(len(list((self.root / 'exports').glob('*.json'))), 2)

    def test_failed_output_is_removed_without_touching_existing_file(self):
        folder = self.root / 'generated'
        (folder / 'kept.docx').write_bytes(b'original')
        with self.assertRaises(RuntimeError):
            with reserve_output(str(folder), 'kept.docx') as (_, stream):
                stream.write(b'partial')
                raise RuntimeError('simulated failure')
        self.assertEqual([p.name for p in folder.iterdir()], ['kept.docx'])
        self.assertEqual((folder / 'kept.docx').read_bytes(), b'original')

    def test_spreadsheet_formula_like_values_remain_text(self):
        for value in ['=1+1', '+123', '-123', '@SUM(1,2)', '  =1+1', '\t=1+1']:
            record = [{'date': '2026-09-01', 'task': value}]
            workbook = load_workbook(io.BytesIO(exporter.to_xlsx_bytes(record)))
            cell = workbook.active['E2']
            self.assertEqual(cell.value, value)
            self.assertEqual(cell.data_type, 's')
            csv_row = list(csv.DictReader(io.StringIO(exporter.to_csv(record))))[0]
            self.assertEqual(csv_row['task'], "'" + value)
            self.assertEqual(json.loads(exporter.to_json(record))[0]['task'], value)

    def test_editor_returns_canonical_values(self):
        draft = self.make_draft([])
        response = self.client.post(f'/draft/{draft}/save', json={'records': [{'date': '2026-09-01', 'day': 'Friday', 'check_in': '9am'}]})
        self.assertEqual(response.status_code, 200)
        record = response.json['rows'][0]['record']
        self.assertEqual((record['day'], record['check_in']), ('Tuesday', '09:00 AM'))

    def test_csrf_requires_session_token_for_form_and_json(self):
        self.app.config['CSRF_ENABLED'] = True
        self.assertEqual(self.client.post('/daily', data={'date': '2026-09-01'}).status_code, 400)
        self.client.get('/daily')
        with self.client.session_transaction() as session:
            token = session['csrf_token']
        self.assertEqual(self.client.post('/daily', data={'date': '2026-09-01', 'csrf_token': '\u00e9'}).status_code, 400)
        self.assertEqual(self.client.post('/daily', data={'date': '2026-09-01', 'csrf_token': token}).status_code, 302)
        draft = self.make_draft([])
        response = self.client.post(f'/draft/{draft}/save', json={'records': []})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.json['ok'])
        response = self.client.post(f'/draft/{draft}/save', json={'records': []}, headers={'X-CSRF-Token': token})
        self.assertEqual(response.status_code, 200)

    def test_repaired_deep_json_is_handled(self):
        # Decoder depth limits differ across Python versions. Exercise the
        # repair retry's failure path without relying on a particular limit.
        with patch('services.json_processor.json.loads', side_effect=[json.JSONDecodeError('Invalid', '', 0), RecursionError()]):
            result = parse_json_text('```json\n[]\n```')
        self.assertFalse(result['ok'])

    def test_dashboard_progress_counts_only_planned_days(self):
        with self.app.app_context():
            upsert_entries([{'date': '2026-09-01', 'task': 'Planned'}, {'date': '2026-08-01', 'task': 'Outside plan'}], get_schedule())
        response = self.client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'1 of 88', response.data)
        self.assertIn(b'Planned', response.data)


if __name__ == '__main__':
    unittest.main(verbosity=2)
