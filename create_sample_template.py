"""Creates a sample Word logbook to practise with:  sample_templates/Internship_Training_Logbook_Sample.docx

It has two tables (Attendance and Daily Progress) whose date cells are pre-filled for the
first 4 weeks of the internship, following the 5-day / 6-day alternating schedule.

Run:  python create_sample_template.py
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from services.common import day_name, format_date
from services.schedule_generator import Schedule

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_templates")
OUT_FILE = os.path.join(OUT_DIR, "Internship_Training_Logbook_Sample.docx")


def style_header(row):
    for cell in row.cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)


def small(table):
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)


def build(weeks=4):
    os.makedirs(OUT_DIR, exist_ok=True)
    schedule = Schedule()
    days = [d for w in range(1, weeks + 1) for d in schedule.working_days(w)]

    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Internship Training Logbook"
    section.footer.paragraphs[0].text = "Student Industrial Training"
    doc.add_heading("Internship Training Logbook", level=0)

    doc.add_heading("Attendance", level=1)
    t1 = doc.add_table(rows=1, cols=4)
    t1.style = "Table Grid"
    for c, text in zip(t1.rows[0].cells, ["Date", "Day", "Check In", "Check Out"]):
        c.text = text
    for d in days:
        cells = t1.add_row().cells
        cells[0].text = format_date(d, "dd Mon yyyy")
    style_header(t1.rows[0])
    small(t1)

    doc.add_page_break()
    doc.add_heading("Daily Progress", level=1)
    t2 = doc.add_table(rows=1, cols=6)
    t2.style = "Table Grid"
    for c, text in zip(t2.rows[0].cells, ["Date", "Task", "Project", "Domain", "Tools", "Learning"]):
        c.text = text
    for d in days:
        cells = t2.add_row().cells
        cells[0].text = format_date(d, "dd Mon yyyy")
    style_header(t2.rows[0])
    small(t2)

    doc.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    print("Created:", build())
